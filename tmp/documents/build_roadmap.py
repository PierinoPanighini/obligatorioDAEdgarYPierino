from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'documents' / 'Hoja de ruta Batalla Naval.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)

milestones = [
    {
        'n': 1, 'title': 'Organización y dominio', 'dates': '1 al 4 de octubre de 2026',
        'period': '1–4 octubre', 'summary': 'Requisitos, primer modelo conceptual y repositorio',
        'goal': 'Contar con una base de requisitos compartida y un primer modelo conceptual revisado por los dos.',
        'checks': [
            'Tenemos una lista de requisitos funcionales, de diseño y de entrega contrastada con la consigna.',
            'Identificamos los conceptos del dominio, sus relaciones y multiplicidades, sin mezclar clases de interfaz o infraestructura.',
            'El primer diagrama conceptual está guardado en ASTAH y ambos podemos explicar sus relaciones.',
            'El repositorio es privado, ambos tenemos acceso y verificamos la nomenclatura, el acceso del docente y el plazo indicado en la consigna.',
            'Registramos las dudas del enunciado, las decisiones iniciales y los temas a consultar en clase.',
        ],
    },
    {
        'n': 2, 'title': 'Base técnica y arquitectura MVP', 'dates': '5 al 11 de octubre de 2026',
        'period': '5–11 octubre', 'summary': 'Aplicación ejecutable e integración del tablero',
        'goal': 'Comprobar una interacción completa entre navegador, presentador, modelo y tablero.',
        'checks': [
            'Ambos podemos compilar y ejecutar el proyecto Java con Maven y Spring Boot desde las instrucciones compartidas.',
            'La estructura de paquetes distingue modelo, vistas y presentadores, y el acceso al modelo utiliza una fachada.',
            'BattleShipBoard.js y CommandDispatcher.js están integrados y conservan sus archivos originales sin cambios.',
            'Una acción de la vista llega al presentador, solicita una operación al modelo y vuelve como comandos de actualización visual.',
            'La sala de operaciones básica se abre desde la ruta raíz / y permite iniciar el recorrido hacia una batalla.',
            'Probamos con dos navegadores una modificación del modelo compartido que actualiza ambas vistas automáticamente.',
        ],
    },
    {
        'n': 3, 'title': 'Preparar una batalla', 'dates': '12 al 18 de octubre de 2026',
        'period': '12–18 octubre', 'summary': 'CU02 completo y batalla preparada para jugar',
        'goal': 'Crear una batalla válida, desplegar la flota y conservar el avance de la preparación.',
        'checks': [
            'Podemos configurar lugar, táctica, dimensiones y cantidades de cada tipo de barco.',
            'Validamos nombre único o autogenerado, dimensiones entre 6 y 12, cantidades de 0 a 6, flota no vacía y ocupación menor al 25 por ciento.',
            'Al continuar la batalla figura En preparación y conserva el avance al salir y volver a entrar.',
            'Lista de barcos, selección inicial, orientación y leyenda cumplen la consigna.',
            'La ubicación respeta la proa indicada, los límites, la superposición y la separación diagonal.',
            'Al colocar un barco se actualiza la lista; al terminar se indica Batalla preparada.',
            'Librar batalla exige toda la flota posicionada y genera una flota enemiga aleatoria con igual composición y reglas válidas.',
            'Probamos los errores de configuración y ubicación, incluido hacer clic en aguas enemigas; no producen cambios inválidos.',
        ],
    },
    {
        'n': 4, 'title': 'Librar una batalla', 'dates': '19 al 25 de octubre de 2026',
        'period': '19–25 octubre', 'summary': 'CU03 completo y tres tácticas de juego',
        'goal': 'Poder terminar una partida y comprobar el comportamiento del combate y sus tácticas.',
        'checks': [
            'El humano comienza; cada disparo válido se resuelve como agua, impacto o hundimiento, según corresponda.',
            'El disparo humano y la respuesta de la máquina forman una ronda indivisible; si gana el humano, no se ejecuta la respuesta.',
            'Secuencial, Espiral y Caza y ataque funcionan mediante polimorfismo; podemos explicar cómo agregar otra táctica.',
            'Comprobamos recorridos en océanos rectangulares y que la máquina no dispara fuera de límites ni repite celdas.',
            'Las bajas enemigas se registran en orden de hundimiento; ronda, estado y tableros reflejan el resultado correcto.',
            'Podemos pausar y retomar una batalla, consultar una finalizada y revelar los barcos enemigos pendientes cuando se pierde.',
            'Disparar a aguas propias, repetir un disparo o intentar actuar en un estado inválido informa el error y conserva el estado válido.',
            'Con dos navegadores en la misma batalla, comprobamos que prevalece el primer disparo registrado de la ronda y que una solicitud atrasada no se aplica como otro turno.',
        ],
    },
    {
        'n': 5, 'title': 'Aplicación completa y sincronización', 'dates': '26 de octubre al 1 de noviembre de 2026',
        'period': '26 octubre–1 noviembre', 'summary': 'CU01 completo, precargas y actualización compartida',
        'goal': 'Completar los tres casos de uso como una aplicación coherente y compartida entre sesiones.',
        'checks': [
            'La sala de operaciones muestra ganadas, perdidas, éxito de campaña y todas las columnas y acciones requeridas.',
            'Precargamos los cinco tipos de barcos, las tres tácticas y al menos una batalla válida en cada estado, con la flota exigida.',
            'Desde la sala podemos crear o entrar a cualquier batalla y regresar conservando el estado correspondiente.',
            'Las sesiones comparten las batallas y reciben los cambios de preparación, combate y estadísticas sin actualizar manualmente ni recargar la página.',
            'La aplicación permite acceso anónimo desde distintos navegadores; cada vista mantiene su contexto de interacción.',
            'Observador y los eventos del modelo propagan las actualizaciones a las vistas interesadas; ambos podemos explicar ese recorrido.',
            'Los controles permanecen habilitados y permiten probar todos los cursos alternativos con sus mensajes de error.',
        ],
    },
    {
        'n': 6, 'title': 'Verificación y cierre del software', 'dates': '2 al 8 de noviembre de 2026',
        'period': '2–8 noviembre', 'summary': 'Software completo para el 9 de noviembre',
        'goal': 'Llegar al 9 de noviembre con las funciones y el diseño implementados y verificados.',
        'checks': [
            'Revisamos los flujos normales y alternativos de los tres casos de uso y registramos los resultados de las pruebas.',
            'Probamos casos límite de configuración, colocación, combate, finalización y acciones simultáneas entre navegadores.',
            'Revisamos MVP con vista pasiva, Experto, Fachada, polimorfismo, Observador, excepciones, paquetes y duplicación de código.',
            'Las validaciones corresponden al modelo y las excepciones se traducen en mensajes; la vista ejecuta las órdenes del presentador.',
            'Una copia limpia del proyecto compila y ejecuta correctamente con las instrucciones compartidas.',
            'No quedan funciones pendientes ni errores que impidan cumplir los casos de uso; identificamos la versión del software lista al 9 de noviembre.',
            'Compartimos una lista concreta de documentación pendiente y ajustes menores para la última semana.',
        ],
    },
    {
        'n': 7, 'title': 'Documentación revisión y margen', 'dates': '9 al 15 de noviembre de 2026',
        'period': '9–15 noviembre', 'summary': 'Documentación final, revisión y paquete de entrega',
        'goal': 'Producir y cerrar la documentación, resolver detalles menores y verificar el paquete final.',
        'checks': [
            'El diagrama conceptual final representa el dominio, con relaciones y multiplicidades revisadas, y abre correctamente en ASTAH.',
            'El diagrama de diseño incluye modelo, vistas y presentadores, refleja la implementación final y está guardado en formato .asta.',
            'El PDF contiene portada con foto, número y nombre de ambos integrantes, grupo y docente.',
            'El PDF incluye la autoevaluación de faltantes, errores o mejoras; si no detectamos ninguno, lo indicamos explícitamente.',
            'Registramos las herramientas de IA utilizadas, el contexto de uso y la verificación del contenido generado.',
            'Ambos revisamos código, diagramas y documento, y podemos explicar toda la solución y el aporte de cada integrante.',
            'Resolvimos los ajustes menores y repetimos las comprobaciones afectadas por esos cambios.',
            'El ZIP contiene el código completo, los diagramas .asta y un único PDF; verificamos el límite de 40 MB.',
            'Descomprimimos el ZIP en una carpeta limpia, ejecutamos el proyecto y comprobamos que abren los documentos incluidos.',
            'Revisamos la inscripción, el equipo en Gestión y el acceso del docente al repositorio; el paquete queda listo para entregar el 16.',
        ],
    },
    {
        'n': 8, 'title': 'Entrega anticipada', 'dates': '16 de noviembre de 2026',
        'period': '16 noviembre', 'summary': 'Entrega realizada y constancia comprobada',
        'goal': 'Entregar el proyecto tres días antes del límite oficial y confirmar la recepción.',
        'checks': [
            'Ambos confirmamos cuál es el ZIP final y que corresponde a la versión revisada.',
            'Subimos el archivo correcto a la instancia de evaluación correspondiente en Gestión.',
            'Comprobamos el nombre, tamaño, fecha y hora registrados en la entrega.',
            'Verificamos la constancia o correo de confirmación y conservamos una copia del archivo entregado.',
        ],
    },
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.75), Cm(1.65)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
sec.footer_distance = Cm(0.8)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Calibri')
    for border in list(style._element.findall('.//' + qn('w:pBdr'))):
        border.getparent().remove(border)

normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Subtitle'].font.italic = False
doc.styles['Subtitle'].paragraph_format.space_after = Pt(14)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)

check_style = doc.styles.add_style('Checklist', 1)
check_style.base_style = normal
check_style.font.size = Pt(10.5)
check_style.paragraph_format.left_indent = Cm(0.6)
check_style.paragraph_format.first_line_indent = Cm(-0.6)
check_style.paragraph_format.tab_stops.add_tab_stop(Cm(0.6))
check_style.paragraph_format.space_after = Pt(4)
check_style.paragraph_format.line_spacing = 1.07
check_style.paragraph_format.keep_together = True

def paragraph(text='', bold_lead=None, size=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    if size:
        for run in p.runs:
            run.font.size = Pt(size)
    return p

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.text = text
        cell.width = Cm(width)
    for record in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, record, widths):
            cell.text = text
            cell.width = Cm(width)
    for row_index, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement('w:cantSplit')
        tr_pr.append(cant_split)
        if row_index == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DDEBF7' if row_index == 0 else 'FFFFFF')
            tc_pr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                val = OxmlElement('w:' + side)
                val.set(qn('w:w'), '90' if side in ['top', 'bottom'] else '110')
                val.set(qn('w:type'), 'dxa')
                margins.append(val)
            tc_pr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement('w:' + side)
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '4')
                border.set(qn('w:color'), 'D9D9D9')
                borders.append(border)
            tc_pr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if col_index == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = row_index == 0
    return t

W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'
check_count = 0
def checklist(text, tag):
    global check_count
    check_count += 1
    p = doc.add_paragraph(style='Checklist')
    sdt = OxmlElement('w:sdt')
    props = OxmlElement('w:sdtPr')
    ident = OxmlElement('w:id')
    ident.set(qn('w:val'), str(2026000 + check_count))
    props.append(ident)
    for element_name in ['alias', 'tag']:
        element = OxmlElement('w:' + element_name)
        element.set(qn('w:val'), tag)
        props.append(element)
    checkbox = OxmlElement('w14:checkbox')
    for name, value in [('checked', '0'), ('checkedState', '2611'), ('uncheckedState', '2610')]:
        element = OxmlElement('w14:' + name)
        element.set('{' + W14 + '}val', value)
        if name != 'checked':
            element.set('{' + W14 + '}font', 'Segoe UI Symbol')
        checkbox.append(element)
    props.append(checkbox)
    sdt.append(props)
    content = OxmlElement('w:sdtContent')
    run = OxmlElement('w:r')
    run_props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Segoe UI Symbol')
    fonts.set(qn('w:hAnsi'), 'Segoe UI Symbol')
    run_props.append(fonts)
    run.append(run_props)
    label = OxmlElement('w:t')
    label.text = '☐'
    run.append(label)
    content.append(run)
    sdt.append(content)
    p._p.append(sdt)
    p.add_run('\t' + text)

doc.add_paragraph('Hoja de ruta de Batalla Naval', 'Title')
doc.add_paragraph('Diseño de aplicaciones con desarrollo generativo', 'Subtitle')
paragraph('Este plan organiza el desarrollo del obligatorio de Edgar y Pierino. El objetivo es tener el software completo el 9 de noviembre y entregar el proyecto el 16 de noviembre de 2026.')
paragraph('La semana del 9 al 15 de noviembre se reserva también para hacer y cerrar la documentación, además de revisar detalles menores y disponer de margen. Los borradores y las decisiones se registran durante el desarrollo.')

doc.add_heading('Fechas objetivo', level=2)
table(['Fecha', 'Resultado esperado'], [
    ('9 noviembre', 'Software completo y verificado'),
    ('9 al 15 noviembre', 'Documentación final, revisión y margen'),
    ('16 noviembre', 'Entrega anticipada confirmada'),
    ('19 noviembre a las 21:00', 'Límite oficial indicado por la consigna, hora de Uruguay'),
], [4.4, 12.6])

doc.add_heading('Equipo y seguimiento', level=2)
paragraph('Integrantes  Edgar y Pierino', size=10)
paragraph('Los dos trabajamos en cada hito y anotamos qué partes hace cada uno. Ambos revisamos la evidencia y debemos comprender la solución completa.', size=10)

doc.add_heading('Resumen de hitos', level=2)
table(['Hito', 'Período', 'Resultado'], [(str(m['n']), m['period'], m['summary']) for m in milestones], [1.1, 4.1, 11.8])

use_p = paragraph('Marcamos una casilla cuando el criterio está comprobado. Cumplimos un hito cuando todas sus casillas están verificadas por ambos. Las fechas son objetivos; las casillas reflejan el avance efectivo.', size=9.5)
use_p.paragraph_format.space_before = Pt(7)

def milestone(m):
    doc.add_heading(f"Hito {m['n']} {m['title']}", level=2)
    date_p = paragraph(m['dates'], size=10)
    date_p.paragraph_format.keep_with_next = True
    goal_p = paragraph(m['goal'], bold_lead='Resultado esperado  ', size=10.5)
    goal_p.paragraph_format.keep_with_next = True
    meta = paragraph('Edgar  Partes del hito ______________________________________________', size=9.5)
    meta.paragraph_format.space_after = Pt(3)
    meta.paragraph_format.keep_with_next = True
    meta2 = paragraph('Pierino  Partes del hito _____________________________________________', size=9.5)
    meta2.paragraph_format.space_after = Pt(8)
    meta2.paragraph_format.keep_with_next = True
    for i, text in enumerate(m['checks'], 1):
        checklist(text, f"Hito {m['n']} Criterio {i}")
    note = paragraph('Evidencia o pendientes _____________________________________________', size=9)
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(8)

for page_title, indexes in [
    ('Organización y base técnica', [0, 1]),
    ('Preparación y combate', [2, 3]),
    ('Integración y cierre del software', [4, 5]),
    ('Documentación y entrega', [6, 7]),
]:
    doc.add_page_break()
    doc.add_heading(page_title, level=1)
    for idx in indexes:
        milestone(milestones[idx])

paragraph('Referencia de requisitos  Obligatorio DADG SEPTIEMBRE 2026, especialmente las páginas 12 y 13 sobre entrega y diseño. Las aclaraciones acordadas en clase se incorporan a las comprobaciones del equipo.', size=9)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Página ').font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

props = doc.core_properties
props.title = 'Hoja de ruta de Batalla Naval'
props.subject = 'Plan de desarrollo y checklist de hitos para dos integrantes'
props.author = 'Edgar y Pierino'
props.keywords = 'Batalla Naval, hoja de ruta, MVP, hitos, documentación'
props.comments = ''
doc.save(OUT)
print(f'Archivo: {OUT}')
print(f'Hitos: {len(milestones)}; casillas: {check_count}')
